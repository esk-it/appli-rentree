"""Un seul export Charlemagne, et tout ce qu'il sait.

L'écran des cartes réclamait un second export — CardStudio — pour
apprendre le code niveau et le code établissement de chaque classe. Et les
mots de passe réseau, que la direction lit dans Charlemagne, ne se
retrouvaient qu'en rouvrant KoXo. L'export de base peut porter tout cela :
il suffit d'y ajouter les colonnes.
"""
from __future__ import annotations

import base64

import openpyxl
import pandas as pd
import pytest

MAITRE = "sardine-clavier-molybdene-1789"


# ---------------------------------------------------------------------------
# La lecture : un mot de passe reste du texte
# ---------------------------------------------------------------------------


def test_les_colonnes_du_compte_reseau_sont_reconnues_et_lues_en_texte(tmp_path):
    """`0123` lu comme un nombre deviendrait `123.0` : un mot de passe faux,
    qu'on croirait bon."""
    from backend.services.parser_charlemagne import lire_htm

    f = tmp_path / "export.htm"
    f.write_text(
        "<table><tr><th>Identifiant élève</th><th>Nom</th><th>Prénom</th>"
        "<th>ID Réseau Péda</th><th>MDP Réseau Péda</th></tr>"
        "<tr><td>8761</td><td>SAILLOUR</td><td>Aaron</td><td>asaillour2</td><td>0123</td></tr>"
        "<tr><td>8762</td><td>MARTIN</td><td>Léa</td><td>lmartin</td><td>4567</td></tr>"
        "</table>",
        encoding="cp1252",
    )
    df = lire_htm(f)
    assert list(df["login_charlemagne"]) == ["asaillour2", "lmartin"]
    assert list(df["mdp_charlemagne"]) == ["0123", "4567"]


def test_un_classeur_garde_aussi_ses_mots_de_passe_intacts(tmp_path):
    from backend.services.parser_charlemagne import lire_xlsx

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Le 24 septembre 2026 à 16 h 02"])
    ws.append([])
    ws.append(["Identifiant élève", "Nom", "Prénom", "MDP Réseau Péda"])
    ws.append([8761, "SAILLOUR", "Aaron", "0456"])
    ws.append([8762, "MARTIN", "Léa", 7890])
    chemin = tmp_path / "export.xlsx"
    wb.save(chemin)

    df = lire_xlsx(chemin)
    assert list(df["mdp_charlemagne"]) == ["0456", "7890"]
    assert list(df["id_charlemagne"]) == [8761, 8762]


# ---------------------------------------------------------------------------
# Les codes de classe, appris à l'ingestion
# ---------------------------------------------------------------------------


@pytest.fixture()
def classes(session, site_factory):
    from backend.models import TableCorrespondance

    ndk, su = site_factory("NDK"), site_factory("SU")
    t = {}
    for site, code in ((su, "31"), (ndk, "2_1"), (ndk, "2_BPMRC1")):
        t[code] = TableCorrespondance(
            site_id=site.id, classe_charlemagne_long=code, classe_code_court=code,
            ou_pre_rentree=f"/{site.nom}", ou_definitive=f"/{site.nom}/{code}",
        )
        session.add(t[code])
    session.commit()
    return t


def _ligne(id_ch, nom, prenom, classe, **kw):
    return {"id_charlemagne": id_ch, "nom": nom, "prenom": prenom,
            "code_classe": classe, **kw}


def _ingerer(session, libelle, lignes, *, mode="reel", cle=None):
    from backend.services.ingestion import RapportIngestion, _ingerer_eleves

    rapport = RapportIngestion(type_personne="eleve", annee_libelle=libelle, mode=mode)
    return _ingerer_eleves(
        session, pd.DataFrame(lignes), libelle, mode, rapport, cle_coffre=cle
    )


def test_l_ingestion_apprend_les_codes_que_cardstudio_reclame(session, classes):
    from backend.models import Personne

    r = _ingerer(session, "2026-2027", [
        _ligne(9001, "A", "Un", "31", code_niveau="3EMES", code_etablissement="02-COL"),
        _ligne(9002, "B", "Deux", "2_1", code_niveau="1-2NDES-LY", code_etablissement="03-LY"),
        _ligne(9003, "C", "Trois", "2_BPMRC1", code_niveau="1-2NDES-LP", code_etablissement="04-LP"),
    ])

    session.expire_all()
    assert (classes["31"].code_niveau, classes["31"].code_etablissement) == ("3EMES", "02-COL")
    assert (classes["2_BPMRC1"].code_niveau, classes["2_BPMRC1"].code_etablissement) == (
        "1-2NDES-LP", "04-LP",
    )
    assert session.query(Personne).filter_by(id_charlemagne=9002).one().niveau == "1-2NDES-LY"
    assert any("3 classe(s) apprise(s)" in a and "3 classe(s) sur 3" in a for a in r.appris)
    assert r.avertissements == []


def test_une_annee_passee_complete_sans_reecrire(session, classes):
    """Une classe peut changer de niveau d'une année à l'autre : l'export
    de l'an dernier ne réécrit pas ce que celui de cette année a appris."""
    _ingerer(session, "2026-2027", [
        _ligne(9001, "A", "Un", "31", code_niveau="3EMES", code_etablissement="02-COL"),
    ])
    _ingerer(session, "2025-2026", [
        _ligne(9001, "A", "Un", "31", code_niveau="VIEUX", code_etablissement="02-COL"),
        _ligne(9002, "B", "Deux", "2_1", code_niveau="1-2NDES-LY", code_etablissement="03-LY"),
    ])

    session.expire_all()
    assert classes["31"].code_niveau == "3EMES"
    assert classes["2_1"].code_niveau == "1-2NDES-LY", "ce qui manquait est complété"


def test_deux_codes_pour_une_classe_se_signalent(session, classes):
    r = _ingerer(session, "2026-2027", [
        _ligne(9001, "A", "Un", "31", code_niveau="3EMES", code_etablissement="02-COL"),
        _ligne(9002, "B", "Deux", "31", code_niveau="4EMES", code_etablissement="02-COL"),
    ])
    session.expire_all()
    assert classes["31"].code_niveau == "3EMES", "le premier lu est gardé"
    assert any("31" in a and "premier lu" in a for a in r.avertissements)


def test_un_export_sans_les_colonnes_dit_quoi_ajouter(session, classes):
    r = _ingerer(session, "2026-2027", [_ligne(9001, "A", "Un", "31")])
    assert any("Code niveau" in a and "Code établissement" in a for a in r.appris)
    assert r.avertissements == [], "rien n'est faux, il manque deux colonnes"


def test_la_carte_prend_les_codes_appris(session, classes, annee_factory):
    from backend.models import Personne
    from backend.services.cartes_cardstudio import construire_fichier, lister_candidats

    _ingerer(session, "2026-2027", [
        _ligne(9002, "B", "Deux", "2_1", code_niveau="1-2NDES-LY", code_etablissement="03-LY"),
    ])
    p = session.query(Personne).filter_by(id_charlemagne=9002).one()

    candidats = lister_candidats(session)
    assert [c.codes_connus for c in candidats] == [True]
    _, rapport = construire_fichier(session, personne_ids=[p.id])
    assert rapport.classes_sans_codes == []


def test_sans_codes_de_classe_la_carte_reprend_ceux_de_l_eleve(
    session, site_factory, annee_factory, personne_factory
):
    """Un élève ingéré avec les colonnes les porte dans son snapshot : la
    carte les reprend quand la classe ne les a pas."""
    import io

    from backend.models import Snapshot, TableCorrespondance
    from backend.services.cartes_cardstudio import construire_fichier

    site = site_factory("NDK")
    annee = annee_factory("2026-2027")
    session.add(TableCorrespondance(
        site_id=site.id, classe_charlemagne_long="2_9", classe_code_court="2_9",
        ou_pre_rentree="/NDK", ou_definitive="/NDK/2_9",
    ))
    p = personne_factory(site_id=site.id)
    session.add(Snapshot(
        personne_id=p.id, annee_scolaire_id=annee.id, nom=p.nom, prenom=p.prenom,
        classe="2_9", niveau="1-2NDES-LY", code_etablissement="03-LY",
    ))
    session.commit()

    octets, rapport = construire_fichier(session, personne_ids=[p.id], annee_id=annee.id)
    ws = openpyxl.load_workbook(io.BytesIO(octets)).active
    entete, ligne = [c.value for c in ws[1]], [c.value for c in ws[2]]
    fiche = dict(zip(entete, ligne))
    assert (fiche["Code niveau"], fiche["Code établissement"], fiche["Etablissement"]) == (
        "1-2NDES-LY", "03-LY", "KREISKER",
    )
    assert rapport.classes_sans_codes == []


# ---------------------------------------------------------------------------
# Les mots de passe réseau, rangés au coffre
# ---------------------------------------------------------------------------


@pytest.fixture()
def cle(session):
    from backend.services.coffre import initialiser

    return initialiser(session, MAITRE)


def _avec_compte(id_ch, nom, prenom, login, mdp):
    return _ligne(id_ch, nom, prenom, "31", login_charlemagne=login, mdp_charlemagne=mdp)


def test_coffre_ouvert_les_mots_de_passe_sont_ranges_avec_leur_identifiant(
    session, classes, cle
):
    from backend.models import SecretConserve
    from backend.services.coffre import chercher

    r = _ingerer(
        session, "2026-2027",
        [_avec_compte(9001, "SAILLOUR", "Aaron", "asaillour2", "Sohp-0123")],
        cle=cle,
    )

    assert (r.nb_mots_de_passe_lus, r.nb_mots_de_passe_ranges, r.coffre_ferme) == (1, 1, False)
    s = session.query(SecretConserve).one()
    assert (s.cible, s.origine, s.identifiant) == ("charlemagne", "charlemagne", "asaillour2")
    assert b"Sohp-0123" not in s.chiffre, "chiffré, jamais en clair"
    lu = chercher(session, cle, "saillour")[0]
    assert (lu.identifiant, lu.mot_de_passe) == ("asaillour2", "Sohp-0123")
    assert all("Sohp-0123" not in a for a in r.appris + r.avertissements)


def test_le_mot_de_passe_de_koxo_n_est_pas_remplace(session, classes, cle):
    """KoXo tient le compte : celui de Charlemagne se range à côté."""
    from backend.models import Personne
    from backend.services.coffre import chercher, deposer

    _ingerer(session, "2026-2027", [_ligne(9001, "SAILLOUR", "Aaron", "31")])
    p = session.query(Personne).one()
    deposer(session, cle, personne_id=p.id, mot_de_passe="koxo-vrai", site="SU")
    session.commit()

    _ingerer(
        session, "2026-2027",
        [_avec_compte(9001, "SAILLOUR", "Aaron", "asaillour", "charlemagne-ancien")],
        cle=cle,
    )
    lus = chercher(session, cle, "saillour")
    assert [(x.cible, x.mot_de_passe) for x in lus] == [
        ("koxo", "koxo-vrai"), ("charlemagne", "charlemagne-ancien"),
    ]


def test_coffre_ferme_rien_n_est_range_et_c_est_dit(session, classes):
    from backend.models import SecretConserve

    r = _ingerer(
        session, "2026-2027",
        [_avec_compte(9001, "SAILLOUR", "Aaron", "asaillour2", "Sohp-0123")],
    )
    assert r.coffre_ferme and r.nb_mots_de_passe_lus == 1
    assert r.nb_mots_de_passe_ranges == 0
    assert session.query(SecretConserve).count() == 0
    assert any("coffre fermé" in a for a in r.appris)


def test_une_simulation_ne_range_rien(session, classes, cle):
    from backend.models import SecretConserve

    r = _ingerer(
        session, "2026-2027",
        [_avec_compte(9001, "SAILLOUR", "Aaron", "asaillour2", "Sohp-0123")],
        mode="simulation", cle=cle,
    )
    assert session.query(SecretConserve).count() == 0
    assert any("seront rangés" in a for a in r.appris)


def test_le_compte_d_une_ancienne_fiche_ne_remplace_pas_l_actuel(
    session, classes, cle
):
    """L'export NDE de l'an dernier porte le compte de l'ancienne fiche
    d'Aaron : il ne doit pas écraser celui de sa fiche de NDK."""
    from backend.models import Personne, SecretConserve, TableCorrespondance
    from backend.services.fusion import fusionner

    nde = TableCorrespondance(
        site_id=classes["31"].site_id, classe_charlemagne_long="4J",
        classe_code_court="4J", ou_pre_rentree="/NDE", ou_definitive="/NDE/4J",
    )
    session.add(nde)
    session.commit()
    _ingerer(session, "2025-2026", [_ligne(717, "SAILLOUR", "Aaron", "4J")])
    _ingerer(
        session, "2026-2027",
        [_avec_compte(8761, "SAILLOUR", "Aaron", "asaillour2", "actuel")], cle=cle,
    )
    ancien = session.query(Personne).filter_by(id_charlemagne=717).one().id
    actuel = session.query(Personne).filter_by(id_charlemagne=8761).one().id
    fusionner(session, ancien, actuel, mode="reel")

    _ingerer(
        session, "2025-2026",
        [{**_ligne(717, "SAILLOUR", "Aaron", "4J"),
          "login_charlemagne": "asaillour1", "mdp_charlemagne": "ancien"}],
        cle=cle,
    )
    s = session.query(SecretConserve).one()
    assert s.identifiant == "asaillour2"


def test_aucun_mot_de_passe_ne_passe_par_le_journal(tmp_db_path, session, classes):
    """Le chemin complet : écran d'ingestion, coffre ouvert."""
    from fastapi.testclient import TestClient

    from backend.main import app
    from backend.models import Generation
    from backend.routers.coffre import verrouiller_maintenant

    html = (
        "<table><tr><th>Identifiant élève</th><th>Nom</th><th>Prénom</th>"
        "<th>Code classe</th><th>ID Réseau Péda</th><th>MDP Réseau Péda</th></tr>"
        "<tr><td>9001</td><td>SAILLOUR</td><td>Aaron</td><td>31</td>"
        "<td>asaillour2</td><td>Tres-Secret-42</td></tr></table>"
    ).encode("cp1252")
    # La clé vit dans le processus, pas dans la base : un coffre laissé
    # ouvert par un test se retrouverait ouvert pour le suivant.
    verrouiller_maintenant()
    try:
        with TestClient(app) as c:
            assert c.post(
                "/api/coffre/initialiser", json={"mot_de_passe": MAITRE}
            ).status_code == 200
            r = c.post("/api/ingestion/base64", json={
                "fichier_base64": base64.b64encode(html).decode(),
                "nom_fichier": "export.htm", "libelle_annee": "2026-2027",
                "type_personne": "eleve", "mode": "reel",
            })
            assert r.status_code == 200, r.text
            assert r.json()["nb_mots_de_passe_ranges"] == 1
            assert "Tres-Secret-42" not in r.text
            trouves = c.get("/api/coffre/chercher", params={"q": "saillour"}).json()
            assert trouves[0]["identifiant"] == "asaillour2"
    finally:
        verrouiller_maintenant()

    session.expire_all()
    for g in session.query(Generation):
        assert "Tres-Secret-42" not in (g.parametres_json + g.resultat_json)
