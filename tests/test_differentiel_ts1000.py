"""Tests du différentiel TS1000, calculé contre l'état réel de la centrale.

Deux règles commandent ce module, et ce sont elles que les tests protègent :
on ne déplace jamais quelqu'un depuis un groupe d'accès, et un CardId ne se
perd jamais. La première a failli coûter quarante-trois accès en septembre
2026 ; la seconde casse une carte encodée sans bruit.
"""
from __future__ import annotations

import io

import pytest

COLONNES = [
    "Op", "Name", "Id", "CardId", "Group", "Technology", "ActivationDate",
    "Grants", "PIN", "ADA",
]


def _export(lignes, colonnes=None):
    """Un export de la centrale.

    Écrit en `.xlsx` : le module lit les deux formats, et fabriquer un
    `.xls` demanderait une dépendance de plus pour rien.
    """
    import openpyxl

    colonnes = colonnes or COLONNES
    classeur = openpyxl.Workbook()
    feuille = classeur.active
    feuille.title = "Sheet 1"
    feuille.append(colonnes)
    for l in lignes:
        feuille.append([l.get(c, "") for c in colonnes])
    tampon = io.BytesIO()
    classeur.save(tampon)
    return tampon.getvalue()


@pytest.fixture()
def contexte(session, site_factory, annee_factory, personne_factory):
    from backend.models import Snapshot, TableCorrespondance

    site = site_factory("NDK")
    annee = annee_factory("2026-2027")
    for code in ("2_5", "2_8", "T_G1"):
        session.add(
            TableCorrespondance(
                site_id=site.id,
                classe_charlemagne_long=code,
                classe_code_court=code,
                ou_pre_rentree="/NDK",
                ou_definitive=f"/NDK/{code}",
            )
        )
    session.commit()

    def _eleve(nom, classe, *, regime="D", id_charlemagne=None):
        p = personne_factory(
            nom=nom, prenom="X", site_id=site.id, id_charlemagne=id_charlemagne
        )
        session.add(
            Snapshot(
                personne_id=p.id,
                annee_scolaire_id=annee.id,
                nom=nom,
                prenom="X",
                classe=classe,
                regime=regime,
            )
        )
        session.commit()
        return p

    return {"site": site, "annee": annee, "eleve": _eleve}


def _calculer(session, contexte, lignes):
    from backend.services.differentiel_ts1000 import calculer_differentiel

    return calculer_differentiel(
        session, _export(lignes), annee_id=contexte["annee"].id
    )


# ---------------------------------------------------------------------------
# Les quatre familles
# ---------------------------------------------------------------------------


def test_un_inscrit_inconnu_de_la_centrale_est_a_creer(session, contexte):
    p = contexte["eleve"]("NEUF", "2_5")
    r = _calculer(session, contexte, [])

    assert [m.badge for m in r.creations] == [str(p.badge)]
    assert r.creations[0].groupe_vise == "2_5"
    assert r.creations[0].op == "a"


def test_un_eleve_dans_la_mauvaise_classe_se_deplace(session, contexte):
    p = contexte["eleve"]("BOUGE", "2_8")
    r = _calculer(
        session, contexte,
        [{"Id": str(p.badge), "Name": "BOUGE X", "Group": "2_5", "CardId": "E012AB"}],
    )

    assert len(r.deplacements) == 1
    m = r.deplacements[0]
    assert (m.groupe_actuel, m.groupe_vise, m.op) == ("2_5", "2_8", "m")


def test_un_eleve_deja_bien_range_ne_bouge_pas(session, contexte):
    p = contexte["eleve"]("STABLE", "2_5")
    r = _calculer(
        session, contexte, [{"Id": str(p.badge), "Group": "2_5", "CardId": "E1"}]
    )
    assert r.est_vide


def test_un_badge_dans_une_classe_sans_inscrit_est_propose_a_la_suppression(
    session, contexte
):
    r = _calculer(
        session, contexte,
        [{"Id": "99999", "Name": "PARTI X", "Group": "2_5", "CardId": "E9"}],
    )
    assert [m.badge for m in r.suppressions] == ["99999"]
    assert r.suppressions[0].op == "b"
    assert any("preuve de départ" in a for a in r.avertissements)


# ---------------------------------------------------------------------------
# La règle qui a failli coûter quarante-trois accès
# ---------------------------------------------------------------------------


def test_un_groupe_dacces_nest_jamais_quitte(session, contexte):
    """Un AVS, un agent, un interne y est rangé pour ouvrir des portes.

    Le ramener vers sa classe les lui ferme. Le groupe n'est pas dans la
    Table : il ne désigne pas une classe, donc on n'y touche pas.
    """
    p = contexte["eleve"]("AVS", "2_5")
    r = _calculer(
        session, contexte,
        [{"Id": str(p.badge), "Name": "AVS X", "Group": "Menage", "CardId": "E7"}],
    )

    assert r.deplacements == []
    assert [m.badge for m in r.proteges] == [str(p.badge)]
    assert "Menage" in r.proteges[0].motif


def test_un_badge_hors_classe_et_hors_referentiel_nest_pas_supprime(session, contexte):
    """Le personnel n'est pas au référentiel : l'absence ne l'accuse pas."""
    r = _calculer(
        session, contexte,
        [{"Id": "500", "Name": "AGENT X", "Group": "Direction", "CardId": "E5"}],
    )
    assert r.suppressions == []
    assert r.est_vide


def test_sans_table_de_correspondance_le_calcul_se_refuse(session, annee_factory):
    """Sans elle, rien ne distingue une classe d'un groupe d'accès."""
    from backend.services.differentiel_ts1000 import (
        DifferentielImpossible,
        calculer_differentiel,
    )

    annee = annee_factory("2026-2027")
    with pytest.raises(DifferentielImpossible, match="table de correspondance"):
        calculer_differentiel(session, _export([]), annee_id=annee.id)


# ---------------------------------------------------------------------------
# L'internat, appris et non deviné
# ---------------------------------------------------------------------------


def test_le_groupe_dinternat_sapprend_sur_les_camarades(session, contexte):
    place = contexte["eleve"]("DEJA", "2_5", regime="P")
    a_placer = contexte["eleve"]("ENCORE", "2_5", regime="P")

    r = _calculer(
        session, contexte,
        [
            {"Id": str(place.badge), "Group": "Internes 2nde", "CardId": "E1"},
            {"Id": str(a_placer.badge), "Group": "2_5", "CardId": "E2"},
        ],
    )

    assert len(r.internat) == 1
    assert r.internat[0].groupe_vise == "Internes 2nde"
    assert r.groupes_internat == ["Internes 2nde"]


def test_le_repli_passe_par_la_famille_de_classes(session, contexte):
    """Aucun pensionnaire en 2_8, mais on en connaît un en 2_5."""
    place = contexte["eleve"]("DEJA", "2_5", regime="P")
    autre = contexte["eleve"]("AUTRE", "2_8", regime="P")

    r = _calculer(
        session, contexte,
        [
            {"Id": str(place.badge), "Group": "Internes 2nde", "CardId": "E1"},
            {"Id": str(autre.badge), "Group": "2_8", "CardId": "E2"},
        ],
    )
    assert [m.groupe_vise for m in r.internat] == ["Internes 2nde"]


def test_un_pensionnaire_sans_repere_est_signale_pas_devine(session, contexte):
    p = contexte["eleve"]("SEUL", "T_G1", regime="P")
    r = _calculer(
        session, contexte, [{"Id": str(p.badge), "Group": "T_G1", "CardId": "E1"}]
    )

    assert r.internat == []
    assert [m.badge for m in r.indetermines] == [str(p.badge)]
    assert any("aucun camarade" in m.motif for m in r.indetermines)


def test_un_pensionnaire_deja_a_linternat_ne_bouge_pas(session, contexte):
    p = contexte["eleve"]("DEDANS", "2_5", regime="P")
    r = _calculer(
        session, contexte,
        [{"Id": str(p.badge), "Group": "Internes 2nde", "CardId": "E1"}],
    )
    assert r.est_vide


# ---------------------------------------------------------------------------
# Le fichier rendu
# ---------------------------------------------------------------------------


def test_le_cardid_survit_a_la_modification(session, contexte):
    """Une ligne de modification qui le perd fait cesser d'ouvrir la carte."""
    import openpyxl

    from backend.services.differentiel_ts1000 import controler_cardid, ecrire_classeur

    p = contexte["eleve"]("BOUGE", "2_8")
    r = _calculer(
        session, contexte,
        [{"Id": str(p.badge), "Name": "BOUGE X", "Group": "2_5", "CardId": "E012FFFE"}],
    )

    octets = ecrire_classeur(r.colonnes, r.deplacements)
    feuille = openpyxl.load_workbook(io.BytesIO(octets)).active
    entete = [c.value for c in next(feuille.iter_rows(max_row=1))]
    ligne = next(feuille.iter_rows(min_row=2, values_only=True))

    assert entete == r.colonnes
    assert ligne[entete.index("CardId")] == "E012FFFE"
    assert ligne[entete.index("Group")] == "2_8"
    assert ligne[entete.index("Op")] == "m"
    assert controler_cardid(r.colonnes, r.deplacements) == []


def test_le_controle_repere_un_cardid_perdu(session, contexte):
    from backend.services.differentiel_ts1000 import controler_cardid

    p = contexte["eleve"]("BOUGE", "2_8")
    r = _calculer(
        session, contexte,
        [{"Id": str(p.badge), "Group": "2_5", "CardId": "E012FFFE"}],
    )
    # On simule la perte que le contrôle doit attraper.
    r.deplacements[0].ligne["CardId"] = ""
    assert controler_cardid(r.colonnes, r.deplacements) == [str(p.badge)]


def test_une_creation_na_pas_de_carte_et_cest_normal(session, contexte):
    from backend.services.differentiel_ts1000 import controler_cardid

    contexte["eleve"]("NEUF", "2_5")
    r = _calculer(session, contexte, [])
    assert controler_cardid(r.colonnes, r.creations) == []


def test_un_fichier_etranger_se_refuse(session, contexte):
    from backend.services.differentiel_ts1000 import (
        DifferentielImpossible,
        calculer_differentiel,
    )

    with pytest.raises(DifferentielImpossible, match="Op|Id|CardId"):
        calculer_differentiel(
            session, _export([], colonnes=["Autre chose"]),
            annee_id=contexte["annee"].id,
        )
